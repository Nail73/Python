import os
import openpyxl
import json
from PIL import Image
from io import BytesIO
from icrawler.builtin import BingImageCrawler
import time
import shutil

# Настройки
DESKTOP_PATH = os.path.join(os.path.expanduser('~'), 'Desktop')
EXCEL_FILE = os.path.join(DESKTOP_PATH, 'products.xlsx')
OUTPUT_DIR = os.path.join(DESKTOP_PATH, 'product_images')
TARGET_SIZE = (400, 400)
MAX_IMAGES_PER_PRODUCT = 1
CONFIG = {
    'delay': 3,
    'feeder_threads': 2,
    'parser_threads': 2,
    'downloader_threads': 2
}


def load_progress():
    try:
        with open('progress.json', 'r') as f:
            return json.load(f).get('last_row', 1)
    except:
        return 1


def save_progress(row_num):
    with open('progress.json', 'w') as f:
        json.dump({'last_row': row_num}, f)


def process_image(img_path, output_path):
    try:
        with Image.open(img_path) as img:
            # Сохраняем пропорции
            img.thumbnail(TARGET_SIZE, Image.LANCZOS)

            # Создаем новое изображение с белым фоном
            new_img = Image.new("RGBA", TARGET_SIZE, (255, 255, 255, 255))

            # Вычисляем позицию для центрирования
            x = (TARGET_SIZE[0] - img.size[0]) // 2
            y = (TARGET_SIZE[1] - img.size[1]) // 2

            # Вставляем изображение по центру
            new_img.paste(img, (x, y))

            # Сохраняем в PNG
            new_img.save(output_path, "PNG", quality=95)

        return True
    except Exception as e:
        print(f"Ошибка обработки изображения: {str(e)}")
        return False


def download_images(product_name, code, output_dir):
    temp_dir = os.path.join(output_dir, f"temp_{code}")
    try:
        os.makedirs(temp_dir, exist_ok=True)

        crawler = BingImageCrawler(
            feeder_threads=CONFIG['feeder_threads'],
            parser_threads=CONFIG['parser_threads'],
            downloader_threads=CONFIG['downloader_threads'],
            storage={'root_dir': temp_dir},
            log_level='ERROR'
        )

        crawler.crawl(
            keyword=product_name,
            max_num=MAX_IMAGES_PER_PRODUCT,
            min_size=(200, 200)
        )

        # Обработка найденных изображений
        for filename in os.listdir(temp_dir):
            if filename.lower().endswith(('.jpg', '.jpeg', '.png')):
                temp_path = os.path.join(temp_dir, filename)
                output_path = os.path.join(output_dir, f"{code}.png")

                if process_image(temp_path, output_path):
                    return True
                break

        return False

    except Exception as e:
        print(f"Ошибка загрузки для {code}: {str(e)}")
        return False
    finally:
        # Удаляем временную директорию с несколькими попытками
        max_attempts = 3
        for attempt in range(max_attempts):
            try:
                if os.path.exists(temp_dir):
                    shutil.rmtree(temp_dir, ignore_errors=True)
                break
            except Exception as e:
                if attempt == max_attempts - 1:
                    print(f"Не удалось удалить временную директорию {temp_dir}: {str(e)}")
                time.sleep(1)


def main():
    print("=== Парсер изображений (Bing) ===")
    print(f"Формат: PNG {TARGET_SIZE[0]}x{TARGET_SIZE[1]}")
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    start_from = load_progress()
    existing_files = set(os.listdir(OUTPUT_DIR))

    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        total = ws.max_row - 1

        processed = 0
        for i, row in enumerate(ws.iter_rows(min_row=start_from, values_only=True), start_from):
            if not row or len(row) < 2 or not row[0] or not row[1]:
                continue

            code, product_name = str(row[0]).strip(), str(row[1]).strip()
            output_file = f"{code}.png"

            if output_file in existing_files:
                print(f"⏩ [{i}/{total}] Пропускаем {code}")
                continue

            if download_images(product_name, code, OUTPUT_DIR):
                print(f"✅ [{i}/{total}] Успешно: {code}.png")
                processed += 1
            else:
                print(f"❌ [{i}/{total}] Ошибка: {code}")

            if i % 10 == 0:
                save_progress(i)
                print(f"Прогресс: {i}/{total}")

    except KeyboardInterrupt:
        print("\nПрервано пользователем")
    except Exception as e:
        print(f"\nКритическая ошибка: {str(e)}")
    finally:
        save_progress(i if 'i' in locals() else 1)
        print(f"\nЗавершено! Обработано: {processed} товаров")


if __name__ == "__main__":
    main()